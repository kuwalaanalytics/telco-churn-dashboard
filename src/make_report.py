"""Step 6: build the client report as a Word document and a PDF.

Run from the project root:  python src/make_report.py
Reads data/clean/findings.json, data/clean/data_profile.json and portfolio/images/,
then writes reports/Telecom-Churn-Retention-Analysis-Report.docx and .pdf.
The PDF is converted with LibreOffice (the `soffice` command must be installed).
"""
from datetime import date
from pathlib import Path
import json
import subprocess

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / "portfolio" / "images"
OUT_DIR = ROOT / "reports"
NAME = "Telecom-Churn-Retention-Analysis-Report"
TITLE = "Telecom Customer Churn and Retention Analysis"
LIVE = "https://kuwalaanalytics.github.io/telecom-churn-retention-analysis/"

F = json.loads((ROOT / "data" / "clean" / "findings.json").read_text())
P = json.loads((ROOT / "data" / "clean" / "data_profile.json").read_text())
B = {c: {g["group"]: g for g in gs} for c, gs in F["breakdowns"].items()}
SVC = {s["service"]: s for s in F["services"]}
O = F["overall"]

NAVY, INK, MUTED = RGBColor(0x14, 0x21, 0x3D), RGBColor(0x14, 0x1C, 0x33), RGBColor(0x6B, 0x75, 0x88)
ORANGE = RGBColor(0xE8, 0x59, 0x0C)
FONT = "Arial"


def pct(v, d=1):
    return f"{v * 100:.{d}f}%"


def money(v):
    return f"${v:,.0f}"


# ---------- document helpers ----------
doc = Document()


def setup():
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)  # A4
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin, sec.bottom_margin = Inches(0.9), Inches(0.8)
    st = doc.styles["Normal"]
    st.font.name, st.font.size, st.font.color.rgb = FONT, Pt(10.5), INK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15
    for name, size, color, before in (("Heading 1", 18, NAVY, 0), ("Heading 2", 13, NAVY, 12), ("Title", 28, NAVY, 0)):
        h = doc.styles[name]
        h.font.name, h.font.size, h.font.bold, h.font.color.rgb = FONT, Pt(size), True, color
        h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        h.element.rPr.rFonts.set(qn("w:asciiTheme"), FONT)
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(before), Pt(6)
        h.paragraph_format.keep_with_next = True


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def para(text="", bold=False, italic=False, color=None, size=None, align=None, after=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def rich(parts, style=None):
    """Paragraph from (text, bold) pairs."""
    p = doc.add_paragraph(style=style)
    for text, bold in parts:
        p.add_run(text).bold = bold
    return p


def bullet(lead, text):
    return rich([(lead, True), (text, False)], style="List Bullet")


def h1(text):
    # Start each section on a new page without an empty paragraph that could leave a blank page.
    doc.add_heading(text, level=1).paragraph_format.page_break_before = True


def h2(text):
    doc.add_heading(text, level=2)


def table(headers, rows, widths, num_from=1):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style, t.alignment, t.autofit = "Table Grid", WD_TABLE_ALIGNMENT.CENTER, False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(9.5), RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "14213D")
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(9.5)
            if i >= num_from:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if n % 2:
                shade(cells[i], "F3F5FA")
    for i, w in enumerate(widths):
        t.columns[i].width = Inches(w)
    last = len(t.rows) - 1
    for n, row in enumerate(t.rows):
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
            pf = row.cells[i].paragraphs[0].paragraph_format
            pf.space_after = pf.space_before = Pt(2)
            pf.keep_with_next = n < last  # keep the table on one page where it fits
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(files, width, caption):
    """One image, or two side by side, centred with a caption."""
    if len(files) > 1:  # side by side in a borderless table
        t = doc.add_table(rows=1, cols=len(files))
        t.alignment, t.autofit = WD_TABLE_ALIGNMENT.CENTER, False
        for cell, f in zip(t.rows[0].cells, files):
            cell.width = Inches(width + 0.1)
            cp = cell.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.keep_with_next = True
            cp.add_run().add_picture(str(IMG / f), width=Inches(width))
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(IMG / files[0]), width=Inches(width))
    para(caption, italic=True, color=MUTED, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)


def footer():
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Kuwala Analytics  |  {TITLE}  |  Page ")
    r.font.size, r.font.color.rgb = Pt(8.5), MUTED
    r2 = p.add_run()
    r2.font.size, r2.font.color.rgb = Pt(8.5), MUTED
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r2._r.append(el)


# ---------- sections ----------
def cover():
    for _ in range(3):
        para()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT / "assets" / "kuwala-logo.jpg"), width=Inches(2.6))
    t = doc.add_paragraph(TITLE, style="Title")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para("Who is leaving, what it costs, and who to call first", size=14, color=MUTED,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=30)
    para("Client Report", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para("Prepared by Kuwala Analytics", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    today = date.today()
    para(f"{today.day} {today:%B %Y}", size=12, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=60)
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    shade(c, "FFF6D6")
    c.text = ""
    cp = c.paragraphs[0]
    cp.add_run("About this report. ").bold = True
    cp.add_run("This is a demonstration report. It uses the public IBM Telco Customer Churn sample dataset to show "
               "how Kuwala Analytics approaches a churn problem. It is not a real client's data, and the figures are "
               "not results achieved for any client. Amounts are in US dollars, as in the dataset.")


def summary():
    v20 = F["value"][-1]
    h1("1. Executive Summary")
    para("This report answers four questions a telecom or subscription business needs answered: how many customers "
         "are leaving and what it costs, which kinds of customers leave most, which individual customers to contact "
         "first, and what it is worth to keep some of them.")
    table(["Measure", "Result"], [
        ["Customers analysed", f"{O['customers']:,}"],
        ["Churn rate (left in the last month)", f"{pct(O['churn_rate'])} ({O['churned']:,} customers)"],
        ["Retention rate", pct(O["retention_rate"])],
        ["Monthly revenue lost to churn", f"{money(O['monthly_revenue_lost'])} ({pct(O['share_revenue_lost'])} of monthly billing)"],
        ["Average monthly charge", f"${O['avg_monthly']:.2f}"],
        ["Average tenure", f"{O['avg_tenure']:.1f} months"],
        ["Customer lifetime value to date (avg monthly charge x avg tenure)", money(O["clv"])],
    ], [3.9, 2.5])
    h2("What we found")
    bullet("Contract type is the strongest signal. ",
           f"Month-to-month customers churn at {pct(B['Contract']['Month-to-month']['churn_rate'])}, against "
           f"{pct(B['Contract']['Two year']['churn_rate'])} on two year contracts.")
    bullet("The first six months are the danger zone. ",
           f"{pct(B['TenureBand']['0-6 months']['churn_rate'])} of customers in their first 6 months leave.")
    bullet("Leavers are the higher-paying customers. ",
           f"They paid ${O['avg_monthly_churned']:.0f} a month on average, against ${O['avg_monthly_stayed']:.0f} for those who stayed.")
    bullet("Support add-ons go with much lower churn. ",
           f"Internet customers with online security churn at {pct(SVC['OnlineSecurity']['churn_with'])}, against "
           f"{pct(SVC['OnlineSecurity']['churn_without'])} without it.")
    bullet("A prediction model can focus the retention effort. ",
           f"The riskiest 20% of customers contained {pct(v20['share_of_all_churners'], 0)} of all leavers.")
    h2("What we recommend")
    para("Look after new customers in their first six months, move month-to-month customers onto contracts, bundle "
         "online security and tech support, shift electronic check payers to automatic payment, review the fibre "
         "optic offer, and call the highest-risk customers every week. Details are in section 7.")
    figure(["01-overall-churn.png"], 2.6, "Figure 1. Overall churn and the revenue that went with it")


def data_section():
    h1("2. The Data")
    para(f"The dataset holds {P['rows']:,} customers, one row each, with {P['columns']} columns. A churn flag marks "
         "customers who left in the last month.")
    table(["Group", "Columns"], [
        ["Services", "Phone, multiple lines, internet type, online security, online backup, device protection, tech support, streaming TV, streaming movies"],
        ["Account", "Tenure (months), contract type, payment method, paperless billing, monthly charges, total charges"],
        ["Demographics", "Gender, senior citizen, partner, dependents"],
        ["Outcome", "Churn (Yes or No)"],
    ], [1.3, 5.1], num_from=9)
    h2("Data quality")
    table(["Check", "Result"], [
        ["Rows and columns", f"{P['rows']:,} rows, {P['columns']} columns"],
        ["Duplicate rows", str(P["duplicate_rows"])],
        ["Duplicate customer IDs", str(P["duplicate_customer_ids"])],
        ["Missing values", f"{P['blank_strings'].get('TotalCharges', 0)} blank values in total charges; none elsewhere"],
    ], [2.4, 4.0], num_from=9)
    h2("Cleaning decisions")
    bullet("Total charges. ", "The column arrived as text. Its 11 blank values all belong to customers with 0 months "
           "of tenure, meaning new customers not yet billed. They were set to 0 and kept, not dropped.")
    bullet("Senior citizen. ", "Recoded from 0 and 1 to No and Yes, to match the other columns.")
    bullet("Duplicates. ", "None found, so no rows were removed.")
    h2("Features added")
    bullet("Tenure band: ", "0-6, 7-12, 13-24, 25-48 and 49-72 months.")
    bullet("Monthly charge band: ", "under $30, $30-49, $50-69, $70-89 and $90 or more.")
    bullet("Number of services: ", "a count of phone, multiple lines, internet and the six add-ons.")
    bullet("No add-ons flag: ", "customers with none of the six add-ons. Add-on comparisons use internet customers "
           "only, because customers without internet cannot buy them.")


def rate_rows(col, order):
    return [[g, f"{B[col][g]['customers']:,}", pct(B[col][g]["churn_rate"]), money(B[col][g]["revenue_lost"])] for g in order]


def findings():
    heads, w = ["Group", "Customers", "Churn rate", "Monthly revenue lost"], [2.6, 1.2, 1.2, 1.4]
    h1("3. Key Findings")
    h2("3.1 Contract type")
    para(f"Customers on month-to-month plans leave at {pct(B['Contract']['Month-to-month']['churn_rate'])}. On one year "
         f"contracts the rate is {pct(B['Contract']['One year']['churn_rate'])}, and on two year contracts "
         f"{pct(B['Contract']['Two year']['churn_rate'])}.")
    table(heads, rate_rows("Contract", ["Month-to-month", "One year", "Two year"]), w)
    h2("3.2 Tenure")
    para(f"Churn falls steadily the longer a customer stays. Half of all leavers had been customers for "
         f"{O['median_tenure_churned']:.0f} months or less.")
    table(heads, rate_rows("TenureBand", ["0-6 months", "7-12 months", "13-24 months", "25-48 months", "49-72 months"]), w)
    figure(["02-churn-by-contract.png", "03-churn-by-tenure.png"], 3.0,
           "Figures 2 and 3. Churn rate by contract type and by time with the company")
    h2("3.3 Payment method")
    para("Customers paying by electronic check leave far more often than customers on any other payment method.")
    table(heads, rate_rows("PaymentMethod", ["Electronic check", "Mailed check", "Bank transfer (automatic)", "Credit card (automatic)"]), w)
    h2("3.4 Internet service and billing")
    para("Fibre optic customers churn at more than twice the rate of DSL customers and account for most of the "
         "lost revenue. Customers on paperless billing also leave more often.")
    rows = rate_rows("InternetService", ["Fiber optic", "DSL", "No"])
    rows[0][0], rows[2][0] = "Fibre optic internet", "No internet"
    rows[1][0] = "DSL internet"
    pb = rate_rows("PaperlessBilling", ["Yes", "No"])
    pb[0][0], pb[1][0] = "Paperless billing: yes", "Paperless billing: no"
    table(heads, rows + pb, w)
    h2("3.5 Demographics")
    para("Senior citizens, and customers without a partner or dependents, leave more often. Gender makes no difference.")
    demo = []
    for col, label in (("SeniorCitizen", "Senior citizen"), ("Partner", "Partner"), ("Dependents", "Dependents"), ("gender", "Gender")):
        for g in (("Yes", "No") if col != "gender" else ("Female", "Male")):
            r = rate_rows(col, [g])[0]
            r[0] = f"{label}: {g.lower() if col != 'gender' else g}"
            demo.append(r)
    table(heads, demo, w)
    h2("3.6 Services linked to lower churn")
    na = F["no_addons_internet"]
    para(f"Among internet customers, online security and tech support go with much lower churn. Streaming services "
         f"make little difference. Customers with no add-ons at all churn at {pct(na['churn_no_addons'])}, against "
         f"{pct(na['churn_some_addons'])} for those with at least one.")
    names = {"OnlineSecurity": "Online security", "TechSupport": "Tech support", "OnlineBackup": "Online backup",
             "DeviceProtection": "Device protection", "StreamingTV": "Streaming TV", "StreamingMovies": "Streaming movies"}
    table(["Service", "Customers who have it", "Churn with it", "Churn without it"],
          [[names[s["service"]], pct(s["adoption"], 0), pct(s["churn_with"]), pct(s["churn_without"])] for s in F["services"]],
          [2.3, 1.6, 1.25, 1.25])
    figure(["04-churn-by-payment-method.png", "05-services-vs-churn.png"], 3.0,
           "Figures 4 and 5. Churn rate by payment method, and with and without each service (internet customers only)")


def segments():
    h1("4. Customer Segments")
    para("Customers were grouped into five named segments using simple rules a sales or service team can apply "
         "without a data scientist. The segments do not overlap, and every customer belongs to exactly one.")
    table(["Segment", "Rule", "Customers", "Churn", "Avg monthly", "Share of leavers"],
          [[s["group"], s["rule"], f"{s['customers']:,}", pct(s["churn_rate"]), f"${s['avg_monthly']:.0f}",
            pct(s["share_churners"])] for s in F["segments"]],
          [1.3, 2.15, 0.95, 0.6, 0.7, 0.75], num_from=2)
    s = {x["group"]: x for x in F["segments"]}
    top2 = s["New and Uncommitted"], s["Fibre Flight Risks"]
    h2("How to read the segments")
    bullet("Highest risk: New and Uncommitted. ",
           f"More than half leave ({pct(top2[0]['churn_rate'])}), and they make up {pct(top2[0]['share_churners'], 0)} of all leavers.")
    bullet("Costliest to lose: Fibre Flight Risks. ",
           f"Established customers paying ${top2[1]['avg_monthly']:.0f} a month, churning at {pct(top2[1]['churn_rate'])}.")
    bullet("Most valuable to protect: Loyal High-Value. ",
           f"They pay ${s['Loyal High-Value']['avg_monthly']:.0f} a month and have paid {money(s['Loyal High-Value']['avg_total_charges'])} "
           f"each to date. Churn is low at {pct(s['Loyal High-Value']['churn_rate'])}, so the aim is to keep it there.")
    bullet("Where to focus: ",
           f"the first two segments are {pct(sum(x['share_customers'] for x in top2), 0)} of customers but "
           f"{pct(sum(x['share_churners'] for x in top2), 0)} of leavers.")
    figure(["06-customer-segments.png"], 4.0, "Figure 6. Churn rate by customer segment")


def model():
    m = F["model"]
    res, best = m["results"], m["chosen"]
    b, base = res[best], F["model_unweighted_baseline"]
    h1("5. Churn Prediction Model")
    para(f"Three models were trained to score each customer's risk of leaving, then tested on {m['test_size']:,} "
         f"customers the models had never seen ({m['test_churners']} of whom left). Leavers are the smaller group, "
         "so the classes were balanced during training to stop the models ignoring them.")
    table(["Model", "Precision", "Recall", "F1", "ROC-AUC"],
          [[k + (" (used for risk scores)" if k == best else ""), f"{v['precision']:.2f}", f"{v['recall']:.2f}",
            f"{v['f1']:.2f}", f"{v['roc_auc']:.3f}"] for k, v in res.items()],
          [3.0, 0.85, 0.85, 0.85, 0.85])
    h2("What the measures mean")
    bullet(f"Recall of {b['recall'] * 100:.0f}%: ", f"of every 100 customers who really left, the model flagged {b['recall'] * 100:.0f} beforehand.")
    bullet(f"Precision of {b['precision'] * 100:.0f}%: ",
           f"of every 100 customers it flagged, {b['precision'] * 100:.0f} really left. The rest received attention they "
           "may not have needed, a small cost compared with losing a customer.")
    bullet(f"F1 of {b['f1']:.2f}: ", "a single score that balances precision and recall.")
    bullet(f"ROC-AUC of {b['roc_auc']:.2f}: ",
           f"pick one leaver and one stayer at random, and the model ranks the leaver as riskier {b['roc_auc'] * 100:.0f} times out of 100.")
    para(f"All three models performed almost identically, so the result does not depend on one technique. {best} was "
         f"used for risk scores because it had the highest cross-validated ROC-AUC ({b['cv_roc_auc']:.3f}), by a small "
         f"margin. Without balancing the classes, recall fell from {res['Logistic Regression']['recall'] * 100:.0f}% to "
         f"{base['recall'] * 100:.0f}%, which would mean missing nearly half the leavers.")
    h2("Risk bands")
    para("Each customer receives a risk score from 0 to 100. The bands below show how well the score matched what actually happened.")
    rb = {g["group"]: g for g in F["risk_bands"]}
    table(["Risk band", "Score range", "Customers", "Actual churn rate"],
          [[k, rng, f"{rb[k]['customers']:,}", pct(rb[k]["churn_rate"])] for k, rng in
           (("High", "60 to 100"), ("Medium", "30 to 59"), ("Low", "0 to 29"))], [1.6, 1.6, 1.6, 1.6])
    h2("Top churn drivers")
    para("The factors the model relies on most are contract type, tenure, total charges to date, internet service "
         "type, online security, tech support and payment method.")
    figure(["07-top-churn-drivers.png"], 4.2, "Figure 7. Top churn drivers (strongest = 100)")


def value():
    val = F["value"]
    v20 = val[-1]
    h1("6. Value of Acting on the Top At-Risk Customers")
    para("Ranking customers by risk score and contacting only the top of the list concentrates effort where it counts.")
    table(["Customers contacted", "How many", "Actual leavers among them", "Share of all leavers reached", "Better than random"],
          [[f"Top {v['top_pct'] * 100:.0f}% by risk", f"{v['customers_contacted']:,}", pct(v["hit_rate"], 0),
            pct(v["share_of_all_churners"], 0), f"{v['lift_vs_random']:.1f}x"] for v in val],
          [1.6, 0.9, 1.4, 1.4, 1.1])
    para(f"Calling customers at random would reach a leaver {pct(O['churn_rate'], 0)} of the time. The risk list "
         f"reaches one {pct(v20['hit_rate'], 0)} to {pct(val[0]['hit_rate'], 0)} of the time.")
    h2("Revenue protected")
    para(f"The leavers in the top 20% were paying {money(v20['monthly_revenue_of_caught'])} a month between them. "
         "What that is worth depends on how many a retention offer keeps.")
    table(["If the offer keeps", "Monthly revenue protected", "Over 12 months"],
          [[f"{s}% of them", money(v20[f"saved_monthly_{s}"]), money(v20[f"saved_annual_{s}"])] for s in (20, 30, 40)],
          [2.2, 2.1, 2.1])
    figure(["08-value-of-targeting.png"], 3.6, "Figure 8. Share of all leavers found among the highest-risk customers")
    h2("How to read these figures")
    para("These are scenarios, not measured outcomes. They come from a back-test on customers who had already left. "
         "The save rates are assumptions, the 12 month figure assumes saved customers stay for the year, and the cost "
         "of the retention offer is not included. A real engagement would test the offer on a small group first and "
         "measure the true save rate.")


def recommendations():
    na = F["no_addons_internet"]
    h1("7. Recommendations")
    recs = [
        ("Win the first six months",
         f"{pct(B['TenureBand']['0-6 months']['churn_rate'])} of customers in their first 6 months leave.",
         "Add a welcome call, a 30 day check-in and a 90 day review for every new customer."),
        ("Move month-to-month customers onto contracts",
         f"Month-to-month churn is {pct(B['Contract']['Month-to-month']['churn_rate'])}, against "
         f"{pct(B['Contract']['One year']['churn_rate'])} on one year contracts.",
         "Offer a discount or a free add-on for committing to 12 months."),
        ("Bundle online security and tech support",
         f"Internet customers with no add-ons churn at {pct(na['churn_no_addons'])}; those with online security churn at "
         f"{pct(SVC['OnlineSecurity']['churn_with'])}.",
         "Trial these services free for 3 months with customers who have no add-ons."),
        ("Shift electronic check payers to automatic payment",
         f"Electronic check payers churn at {pct(B['PaymentMethod']['Electronic check']['churn_rate'])}, about three times "
         "the rate of customers on automatic payments.",
         "Give a small bill credit for switching."),
        ("Review the fibre optic experience and price",
         f"Fibre customers churn at {pct(B['InternetService']['Fiber optic']['churn_rate'])} and account for "
         f"{money(B['InternetService']['Fiber optic']['revenue_lost'])} of the {money(O['monthly_revenue_lost'])} lost each month.",
         "Survey leavers and compare price and service quality with competitors."),
        ("Call the highest-risk customers every week",
         f"The riskiest 20% of customers contained {pct(F['value'][-1]['share_of_all_churners'], 0)} of all leavers.",
         "Work the at-risk list in the dashboard from the top, starting with high-value customers."),
    ]
    table(["#", "Action", "Evidence", "First step"],
          [[i + 1, a, e, s] for i, (a, e, s) in enumerate(recs)], [0.3, 1.7, 2.3, 2.1], num_from=9)
    h2("A caution")
    para("These patterns show what goes together with churn. They are not proof of cause. For example, customers who "
         "choose a two year contract may already be more loyal. Each recommendation should be run as a small test "
         "first and measured before it is rolled out.")


def dashboard():
    h1("8. The Interactive Dashboard")
    para("All of the analysis is available in one interactive dashboard. It has filters for contract, internet "
         "service, payment method, tenure band and senior citizen, seven key figures, six charts, a sortable and "
         "searchable list of at-risk customers, the model results and the recommended actions. It works on a laptop "
         "or phone, and offline.")
    rich([("Open it here: ", True), (LIVE, False)])
    figure(["dashboard-screenshot.png"], 5.2, "Figure 9. The dashboard with no filters applied")


def method():
    h1("9. Method and Limitations")
    h2("Method")
    for lead, text in [
        ("Clean and prepare. ", "Fixed data types, handled missing values and added the features described in section 2."),
        ("Explore. ", "Measured churn rate and lost revenue across every customer attribute."),
        ("Segment. ", "Grouped customers with simple, non-overlapping rules."),
        ("Predict. ", "Compared logistic regression, random forest and gradient boosting on a 20% holdout, with balanced class weights."),
        ("Score. ", "Risk scores are out-of-fold predictions from 5-fold cross-validation, so no customer is scored by a model trained on their own record."),
        ("Value. ", "Ranked customers by risk and measured how many real leavers sat at the top of the list."),
    ]:
        bullet(lead, text)
    h2("Limitations")
    for lead, text in [
        ("Sample data. ", "This is the public IBM sample dataset, not a real client's data."),
        ("One snapshot. ", "The data covers a single month of churn, so seasonal patterns and trends cannot be seen."),
        ("Association, not cause. ", "The findings show which groups leave more often, not why."),
        ("No cost data. ", "The cost of serving customers and of retention offers is not in the data, so figures are revenue, not profit."),
        ("No reasons for leaving. ", "Complaints, service quality and competitor offers are not recorded."),
    ]:
        bullet(lead, text)
    h2("Tools")
    para("Python (pandas, scikit-learn, Plotly), Chart.js, HTML, CSS and JavaScript.")


def main():
    setup()
    footer()
    cover()
    for section in (summary, data_section, findings, segments, model, value, recommendations, dashboard, method):
        section()
    OUT_DIR.mkdir(exist_ok=True)
    docx_path = OUT_DIR / f"{NAME}.docx"
    doc.save(docx_path)
    print("Wrote", docx_path)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT_DIR), str(docx_path)],
                   check=True, capture_output=True)
    print("Wrote", OUT_DIR / f"{NAME}.pdf")


if __name__ == "__main__":
    main()
